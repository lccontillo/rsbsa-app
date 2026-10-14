"""Monthly accomplishment report generation (DOCX -> PDF via LibreOffice).

Used for COA/LGU-style reporting: certificates issued per barangay and per
secretary for a given month, with the full detail list.
"""
import os
import shutil
import tempfile

from docx import Document
from docx.shared import Pt

from docgen import docx_to_pdf

MONTH_NAMES = ["January", "February", "March", "April", "May", "June",
               "July", "August", "September", "October", "November", "December"]
SHORT_MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
                "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def _fmt_dt(iso_ts: str) -> str:
    """'2026-10-13T06:05:00+08:00' -> 'Oct 13, 2026 06:05 AM'."""
    try:
        y, m, d = int(iso_ts[0:4]), int(iso_ts[5:7]), int(iso_ts[8:10])
        hh, mm = int(iso_ts[11:13]), int(iso_ts[14:16])
        ampm = "AM" if hh < 12 else "PM"
        h12 = hh % 12 or 12
        return f"{SHORT_MONTHS[m - 1]} {d}, {y} {h12:02d}:{mm:02d} {ampm}"
    except Exception:
        return iso_ts or ""


def generate_accomplishment_pdf(payload: dict) -> str:
    """Build the report DOCX in a temp dir, convert to PDF, return PDF path.

    The caller must call cleanup_pdf() when the file has been sent.
    """
    tmpdir = tempfile.mkdtemp(prefix="rsbsa_report_")
    try:
        docx_path = os.path.join(tmpdir, "report.docx")
        _build_docx(docx_path, payload)
        return docx_to_pdf(docx_path, tmpdir)
    except Exception:
        shutil.rmtree(tmpdir, ignore_errors=True)
        raise


def cleanup_pdf(pdf_path: str) -> None:
    """Remove the temp directory that holds the generated report."""
    d = os.path.dirname(os.path.abspath(pdf_path))
    if os.path.basename(d).startswith("rsbsa_report_"):
        shutil.rmtree(d, ignore_errors=True)


def _build_docx(path: str, p: dict) -> None:
    doc = Document()

    def center(text, bold=False, size=None, italic=False):
        para = doc.add_paragraph()
        para.alignment = 1
        run = para.add_run(text)
        run.bold = bold
        run.italic = italic
        if size:
            run.font.size = Pt(size)
        return para

    def line(text, bold=False, size=None):
        para = doc.add_paragraph()
        run = para.add_run(text)
        run.bold = bold
        if size:
            run.font.size = Pt(size)
        return para

    # ---------------- header ----------------
    if p.get("province"):
        center(f"PROVINCE OF {p['province'].upper()}", size=11)
    if p.get("municipality"):
        center(p["municipality"].upper(), bold=True, size=12)
    if p.get("barangay_label"):
        center(p["barangay_label"].upper(), bold=True, size=12)
    center("MONTHLY ACCOMPLISHMENT REPORT", bold=True, size=14)
    center("Registry System for Basic Sectors in Agriculture (RSBSA) — Certification Issuance",
           size=10, italic=True)
    doc.add_paragraph()

    line(f"Period: {p['month_label']}", bold=True)
    line(f"Scope: {p['scope_label']}")
    line(f"Generated: {p['generated_at']} by {p['prepared_by']} (system-generated)")
    doc.add_paragraph()

    # ---------------- summary by barangay ----------------
    line("I. SUMMARY BY BARANGAY", bold=True)
    t = doc.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    hdr[0].text = "Barangay"
    hdr[1].text = "Certificates Issued"
    for c in hdr:
        for para in c.paragraphs:
            for run in para.runs:
                run.bold = True
    for item in p["by_barangay"]:
        row = t.add_row().cells
        row[0].text = str(item["name"])
        row[1].text = str(item["count"])
    total = t.add_row().cells
    total[0].text = "TOTAL"
    total[1].text = str(p["totals"])
    for cell in total:
        for para in cell.paragraphs:
            for run in para.runs:
                run.bold = True
    doc.add_paragraph()

    # ---------------- by secretary ----------------
    line("II. SUMMARY BY CERTIFYING OFFICER", bold=True)
    t2 = doc.add_table(rows=1, cols=3)
    t2.style = "Table Grid"
    hdr = t2.rows[0].cells
    for i, h in enumerate(["Certifying Officer", "Barangay", "Certificates"]):
        hdr[i].text = h
        for para in hdr[i].paragraphs:
            for run in para.runs:
                run.bold = True
    for item in p["by_secretary"]:
        row = t2.add_row().cells
        row[0].text = str(item["name"])
        row[1].text = str(item["barangay"])
        row[2].text = str(item["count"])
    if not p["by_secretary"]:
        row = t2.add_row().cells
        row[0].text = "—"
        row[1].text = "—"
        row[2].text = "0"
    doc.add_paragraph()

    # ---------------- detail list ----------------
    line("III. DETAIL OF CERTIFICATES ISSUED", bold=True)
    if p["rows"]:
        t3 = doc.add_table(rows=1, cols=6)
        t3.style = "Table Grid"
        hdr = t3.rows[0].cells
        for i, h in enumerate(["#", "Certificate No.", "Farmer",
                               "Barangay", "Date Issued", "Issued By"]):
            hdr[i].text = h
            for para in hdr[i].paragraphs:
                for run in para.runs:
                    run.bold = True
        for i, r in enumerate(p["rows"], start=1):
            row = t3.add_row().cells
            row[0].text = str(i)
            row[1].text = str(r["cert_no"])
            row[2].text = str(r["farmer_name"] or "—")
            row[3].text = str(r["barangay"])
            row[4].text = _fmt_dt(r["generated_at"])
            row[5].text = str(r["generator"])
    else:
        line("No certificates were issued during this period.", italic=True)

    # ---------------- signatures ----------------
    doc.add_paragraph()
    doc.add_paragraph()
    center(f"Prepared by: {p['prepared_by']}", bold=True)
    center("Certifying Officer")
    doc.add_paragraph()
    line("Noted:")
    doc.add_paragraph()
    line("____________________________________")
    line("Municipal Agriculturist / Office Head")

    doc.save(path)
